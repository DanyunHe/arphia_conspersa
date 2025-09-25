import os
import cv2
import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.utils.dataframe import dataframe_to_rows

def crop_wing_func(
    wing_0_im,
    wing_1_im,
    file_save_crop,
    sorted_file_list_full_0,
    sorted_file_list_full_1,
    file_path_save,
    population_id,
    file_folder
):
    """
    Crop two wing-mask images with a margin and log the crop in an Excel sheet.
    """
    w_0 = cv2.cvtColor(wing_0_im, cv2.COLOR_BGR2GRAY)
    w_1 = cv2.cvtColor(wing_1_im, cv2.COLOR_BGR2GRAY)

    nx0, ny0 = w_0.shape[1], w_0.shape[0]
    nx1, ny1 = w_1.shape[1], w_1.shape[0]

    # ---------- find crop bounds ----------
    def find_min_x(img):
        for i in range(img.shape[1]):
            if not all(x == 255 for x in img[:, i]):
                return i
        return 0

    def find_max_x(img):
        for i in range(img.shape[1] - 1, -1, -1):
            if not all(x == 255 for x in img[:, i]):
                return i
        return img.shape[1] - 1

    def find_min_y(img):
        for i in range(img.shape[0]):
            if not all(x == 255 for x in img[i, :]):
                return i
        return 0

    def find_max_y(img):
        for i in range(img.shape[0] - 1, -1, -1):
            if not all(x == 255 for x in img[i, :]):
                return i
        return img.shape[0] - 1

    min_x_w = find_min_x(w_0)
    max_x_w = find_max_x(w_0)
    min_y_w = find_min_y(w_0)
    max_y_w = find_max_y(w_0)

    # margin of ±50 pixels where possible
    min_x_w = max(min_x_w - 50, 0)
    min_y_w = max(min_y_w - 50, 0)
    max_x_w = min(max_x_w + 50, min(nx0, nx1))
    max_y_w = min(max_y_w + 50, min(ny0, ny1))

    # ---------- crop and save ----------
    w_0_im_crop = wing_0_im[min_y_w:max_y_w, min_x_w:max_x_w]
    w_1_im_crop = wing_1_im[min_y_w:max_y_w, min_x_w:max_x_w]

    cv2.imwrite(
        os.path.join(file_save_crop,
                     os.path.splitext(sorted_file_list_full_0)[0][:-5] + "_crop.tif"),
        w_0_im_crop,
    )
    cv2.imwrite(
        os.path.join(file_save_crop,
                     os.path.splitext(sorted_file_list_full_1)[0][:-5] + "_crop.tif"),
        w_1_im_crop,
    )

    # ---------- record to Excel ----------
    # These variables must be defined by the caller before use:
    #   file_path_save, population_id, file_folder
    wing_category = os.path.splitext(sorted_file_list_full_0)[0][-7:-9]
    append_data = pd.DataFrame([{
        "filename": os.path.splitext(sorted_file_list_full_0)[0][:-5],
        "w0": nx0,
        "h0": ny0,
        "w1": nx1,
        "h1": ny1,
        f"c_ax_{wing_category}": min_x_w,
        f"c_bx_{wing_category}": max_x_w,
        f"c_ay_{wing_category}": min_y_w,
        f"c_by_{wing_category}": max_y_w,
        f"cw_{wing_category}": w_0_im_crop.shape[1],
        f"ch_{wing_category}": w_0_im_crop.shape[0],
        "folder": file_folder,
    }])

    excel_path = os.path.join(
        file_path_save, f"cropping_log_population_{population_id}.xlsx"
    )

    if os.path.isfile(excel_path):
        wb = load_workbook(excel_path)
    else:
        wb = Workbook()
        wb.save(excel_path)
        wb = load_workbook(excel_path)

    sheet_name = f"crop_{wing_category}"
    if sheet_name not in wb.sheetnames:
        ws = wb.create_sheet(sheet_name)
        ws.append([
            "filename", "w0", "h0", "w1", "h1",
            f"c_ax_{wing_category}", f"c_bx_{wing_category}",
            f"c_ay_{wing_category}", f"c_by_{wing_category}",
            f"cw_{wing_category}", f"ch_{wing_category}", "folder"
        ])
    ws = wb[sheet_name]

    for r in dataframe_to_rows(append_data, index=False, header=False):
        ws.append(r)
    wb.save(excel_path)
